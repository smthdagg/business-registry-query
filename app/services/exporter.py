"""导出：把批量查询结果写成 CSV / XLSX / JSON。"""

import csv
import io
import json

from ..sources.base import FIELD_META

# 表头：关键词 / 角度 / 状态列 + 企业字段
_EXTRA_COLS = [
    ("keyword", "关键词"),
    ("angle", "查询角度"),
    ("result", "查询结果"),
    ("note", "备注"),
]
_HEADERS = _EXTRA_COLS + list(FIELD_META)


_FORMULA_LEAD = ("=", "+", "-", "@", "\t", "\r")


def _safe(v: str) -> str:
    """防公式注入：Excel/LibreOffice 会把 =+-@ 开头的单元格当公式执行。"""
    return "'" + v if v.startswith(_FORMULA_LEAD) else v


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        parts = []
        for v in value:
            if isinstance(v, dict) and v.get("name"):
                pct = f"（{v['percent']}）" if v.get("percent") else ""
                parts.append(f"{v['name']}{pct}")
            elif v:
                parts.append(str(v))
        return "、".join(parts)
    if isinstance(value, dict):
        if value.get("name"):
            pct = f"（{value['percent']}）" if value.get("percent") else ""
            return f"{value['name']}{pct}"
        return json.dumps(value, ensure_ascii=False)
    return _safe(str(value))


def build_rows(items: list[dict]) -> list[dict[str, str]]:
    """items -> 每行一个公司的扁平行（全部导出到 xlsx/csv 的行）。"""
    rows: list[dict[str, str]] = []
    for idx, item in enumerate(items, 1):
        best = item.get("best")
        if item.get("ok") and best:
            row = {
                "keyword": item.get("keyword", ""),
                "angle": item.get("angle", ""),
                "result": "命中" if item.get("count") else "无匹配",
                "note": (
                    f"共 {item['count']} 条匹配"
                    + (f"，另有 {item['count'] - 1} 条见任务详情" if item["count"] > 1 else "")
                ),
            }
            for key, _label in FIELD_META:
                row[key] = _cell(best.get(key))
            rows.append(row)
        else:
            rows.append({
                "keyword": item.get("keyword", ""),
                "angle": item.get("angle", ""),
                "result": "失败" if not item.get("ok") else "无匹配",
                "note": item.get("error", ""),
            })
    return rows


def _header_keys() -> list[str]:
    return [k for k, _ in _HEADERS]


def to_csv(items: list[dict]) -> bytes:
    buf = io.StringIO()
    buf.write("\ufeff")  # BOM：Excel 打开不乱码
    writer = csv.writer(buf)
    writer.writerow([label for _k, label in _HEADERS])
    keys = _header_keys()
    for row in build_rows(items):
        writer.writerow([_safe(row.get(k, "")) for k in keys])
    return buf.getvalue().encode("utf-8")


def to_xlsx(items: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "批量查询结果"
    ws.append([label for _k, label in _HEADERS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center")
    keys = _header_keys()
    for row in build_rows(items):
        ws.append([row.get(k, "") for k in keys])
    # 简单列宽
    for col_idx, (key, label) in enumerate(_HEADERS, 1):
        width = max(len(label), 10)
        if key in ("businessScope", "regLocation", "name"):
            width = 40
        elif key in ("keyword",):
            width = 26
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_json(items: list[dict]) -> bytes:
    payload = {"count": len(items), "items": items}
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


FORMATS = {"csv": to_csv, "xlsx": to_xlsx, "json": to_json}
